"""Render the humanised paper (markdown subset) to DOCX and PDF."""
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "paper_humanised.md")
DOCX_OUT = os.path.join(HERE, "XAI_Modality_Bias_Paper_Humanised.docx")
PDF_OUT = os.path.join(HERE, "XAI_Modality_Bias_Paper_Humanised.pdf")

# ---------------------------------------------------------------- parsing


def parse_blocks(text):
    lines = text.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s == "---":
            blocks.append(("hr", ""))
            i += 1
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("eq", " ".join(x for x in buf if x)))
            continue
        if s.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            rows = []
            for r in buf:
                if re.match(r"^\|[\s\-:|]+\|$", r):
                    continue
                cells = [c.strip() for c in r.strip("|").split("|")]
                rows.append(cells)
            blocks.append(("table", rows))
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            blocks.append(("img", (m.group(1), m.group(2))))
            i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,4}\s|>\s|\||!\[|\`\`\`|---$)", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|<sub>.+?</sub>|<sup>.+?</sup>|\*[^*]+?\*)")


def inline_segments(text):
    """Yield (text, {bold, italic, code, sub, sup}) segments."""
    segs = []
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            segs.append((part[2:-2], {"bold": True}))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            segs.append((part[1:-1], {"code": True}))
        elif part.startswith("<sub>"):
            segs.append((part[5:-6], {"sub": True}))
        elif part.startswith("<sup>"):
            segs.append((part[5:-6], {"sup": True}))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            segs.append((part[1:-1], {"italic": True}))
        else:
            segs.append((part, {}))
    return segs


def strip_markup(text):
    out = []
    for seg, flags in inline_segments(text):
        out.append(seg)
    return "".join(out)


# ---------------------------------------------------------------- DOCX

def build_docx(blocks):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")

    def add_runs(par, text, size=None, color=None):
        for seg, flags in inline_segments(text):
            run = par.add_run(seg)
            run.bold = flags.get("bold", False)
            run.italic = flags.get("italic", False)
            if flags.get("code"):
                run.font.name = "Consolas"
                run.font.size = Pt((size or 11) - 1)
            elif size:
                run.font.size = Pt(size)
            if flags.get("sub"):
                run.font.subscript = True
            if flags.get("sup"):
                run.font.superscript = True
            if color is not None:
                run.font.color.rgb = color

    def footer_page_numbers(section):
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        inner = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        inner.append(t)
        fld.append(inner)
        run._r.addnext(fld)

    footer_page_numbers(sec)

    table_no = 0
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(strip_markup(payload))
            r.bold = True
            r.font.size = Pt(17)
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(strip_markup(payload))
            r.bold = True
            r.font.size = Pt(13)
        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(strip_markup(payload))
            r.bold = True
            r.font.size = Pt(11.5)
        elif kind == "h4":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(strip_markup(payload))
            r.bold = True
            r.italic = True
            r.font.size = Pt(11)
        elif kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)
            pPr.append(pbdr)
        elif kind == "p":
            if payload.startswith("**Table") or payload.startswith("**Figure"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.keep_with_next = True
                add_runs(p, payload, size=9.5)
            elif payload.startswith("**Keywords:**"):
                p = doc.add_paragraph()
                add_runs(p, payload, size=10.5)
            else:
                p = doc.add_paragraph()
                add_runs(p, payload)
        elif kind == "eq":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            add_runs(p, payload, size=10.5)
        elif kind == "code":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(payload)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif kind == "img":
            alt, path = payload
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(os.path.join(HERE, path), width=Inches(6.2))
        elif kind == "table":
            table_no += 1
            rows = payload
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            usable = 6.5
            weights = []
            for ci in range(ncols):
                mx = max(len(strip_markup(r[ci])) if ci < len(r) else 1 for r in rows)
                weights.append(max(6.0, min(mx, 40)) ** 0.85)
            tot = sum(weights)
            col_w = [Inches(usable * w / tot) for w in weights]
            for ci, w in enumerate(col_w):
                table.columns[ci].width = w
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncols):
                    cells[ci].width = col_w[ci]
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ""
                    par = cells[ci].paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.space_before = Pt(2)
                    par.paragraph_format.line_spacing = 1.0
                    par.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if ci > 0 and ri > 0 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    add_runs(par, txt, size=8.5)
                    if ri == 0:
                        for run in par.runs:
                            run.bold = True
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:fill"), "ECECEC")
                        cells[ci]._tc.get_or_add_tcPr().append(shd)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc.save(DOCX_OUT)
    print("wrote", DOCX_OUT)


# ---------------------------------------------------------------- PDF

def build_pdf(blocks):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate, Frame, Image, KeepTogether, PageTemplate,
        Paragraph, Spacer, Table, TableStyle, PageBreak,
    )

    fd = "/usr/share/fonts/truetype/dejavu"
    pdfmetrics.registerFont(TTFont("DJSerif", os.path.join(fd, "DejaVuSerif.ttf")))
    pdfmetrics.registerFont(TTFont("DJSerif-Bold", os.path.join(fd, "DejaVuSerif-Bold.ttf")))
    pdfmetrics.registerFontFamily(
        "DJSerif", normal="DJSerif", bold="DJSerif-Bold",
        italic="DJSerif", boldItalic="DJSerif-Bold",
    )
    pdfmetrics.registerFont(TTFont("DJMono", os.path.join(fd, "DejaVuSansMono.ttf")))

    body = ParagraphStyle(
        "body", fontName="DJSerif", fontSize=9.6, leading=14.2,
        alignment=TA_JUSTIFY, spaceAfter=7.0,
    )
    title = ParagraphStyle(
        "title", parent=body, fontName="DJSerif-Bold", fontSize=16.5,
        leading=22, alignment=TA_CENTER, spaceAfter=8,
    )
    meta = ParagraphStyle("meta", parent=body, alignment=TA_CENTER, spaceAfter=1.5, fontSize=9.3)
    h2 = ParagraphStyle(
        "h2", parent=body, fontName="DJSerif-Bold", fontSize=12.6, leading=16,
        alignment=TA_LEFT, spaceBefore=13, spaceAfter=5, keepWithNext=1,
    )
    h3 = ParagraphStyle(
        "h3", parent=body, fontName="DJSerif-Bold", fontSize=10.6, leading=14,
        alignment=TA_LEFT, spaceBefore=10, spaceAfter=3.5, keepWithNext=1,
    )
    h4 = ParagraphStyle(
        "h4", parent=body, fontName="DJSerif-Bold", fontSize=9.9, leading=13,
        alignment=TA_LEFT, spaceBefore=8, spaceAfter=3, keepWithNext=1,
    )
    eq = ParagraphStyle("eq", parent=body, alignment=TA_CENTER, fontSize=9.3, spaceBefore=5, spaceAfter=9)
    cap = ParagraphStyle("cap", parent=body, fontSize=8.6, leading=12.4, spaceBefore=3, spaceAfter=11)
    code = ParagraphStyle(
        "code", parent=body, fontName="DJMono", fontSize=7.8, leading=11,
        alignment=TA_LEFT, leftIndent=14, spaceAfter=10,
    )
    cellh = ParagraphStyle("cellh", parent=body, fontName="DJSerif-Bold", fontSize=7.9, leading=10.6, alignment=TA_LEFT)
    cell = ParagraphStyle("cell", parent=body, fontSize=7.9, leading=10.6, alignment=TA_LEFT)
    cellc = ParagraphStyle("cellc", parent=cell, alignment=TA_CENTER)
    kw = ParagraphStyle("kw", parent=body, fontSize=9.2)
    capcell = ParagraphStyle("capcell", parent=body, fontSize=8.6, leading=12.2,
                             alignment=TA_LEFT, spaceAfter=0)

    def esc(text):
        t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        for tag in ("sub", "sup"):
            t = t.replace("&lt;" + tag + "&gt;", "<" + tag + ">")
            t = t.replace("&lt;/" + tag + "&gt;", "</" + tag + ">")
        return t

    def mk(text, style):
        t = esc(text)
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
        t = re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"<i>\1</i>", t)
        t = re.sub(r"`(.+?)`", r'<font name="DJMono" size="8.2">\1</font>', t)
        return Paragraph(t, style)

    FRAME_W = letter[0] - 2 * inch
    story = []

    # pre-merge image + caption pairs so they never split across pages
    merged = []
    idx = 0
    while idx < len(blocks):
        kind, payload = blocks[idx]
        if (
            kind == "img"
            and idx + 1 < len(blocks)
            and blocks[idx + 1][0] == "p"
            and blocks[idx + 1][1].startswith("**Figure")
        ):
            merged.append(("fig", (payload, blocks[idx + 1][1])))
            idx += 2
            continue
        if (
            kind == "p"
            and payload.startswith("**Table")
            and idx + 1 < len(blocks)
            and blocks[idx + 1][0] == "table"
        ):
            merged.append(("tbl", (payload, blocks[idx + 1][1])))
            idx += 2
            continue
        merged.append((kind, payload))
        idx += 1
    blocks = merged

    for kind, payload in blocks:
        if kind == "fig":
            (alt, path), caption = payload
            from reportlab.lib.utils import ImageReader
            iw, ih = ImageReader(os.path.join(HERE, path)).getSize()
            w = FRAME_W * (0.94 if ("fig3" in path or "fig4" in path) else 1.0)
            h = w * ih / iw
            max_h = 5.2 * inch
            if h > max_h:
                h = max_h
                w = h * iw / ih
            story.append(KeepTogether([
                Spacer(1, 6),
                Image(os.path.join(HERE, path), width=w, height=h),
                mk(caption, cap),
            ]))
        elif kind == "h1":
            story.append(mk(payload, title))
        elif kind == "h2":
            story.append(mk(payload, h2))
        elif kind == "h3":
            story.append(mk(payload, h3))
        elif kind == "h4":
            story.append(mk(payload, h4))
        elif kind == "hr":
            story.append(Spacer(1, 6))
        elif kind == "p":
            if payload.startswith("**Table") or payload.startswith("**Figure"):
                story.append(mk(payload, cap))
            elif payload.startswith("**Keywords:**"):
                story.append(mk(payload, kw))
            elif payload.startswith("**Nishant") or payload.startswith("Project team") \
                    or payload.startswith("Sardar Vallabhbhai") or payload.startswith("Integrated Artificial"):
                story.append(mk(payload, meta))
            else:
                story.append(mk(payload, body))
        elif kind == "eq":
            story.append(mk(payload, eq))
        elif kind == "code":
            lines = payload.split("\n")
            story.append(Paragraph("<br/>".join(esc(l).replace(" ", "&nbsp;") for l in lines), code))
        elif kind == "img":
            alt, path = payload
            from reportlab.lib.utils import ImageReader
            iw, ih = ImageReader(os.path.join(HERE, path)).getSize()
            w = FRAME_W * (0.94 if "fig3" in path or "fig4" in path else 1.0)
            h = w * ih / iw
            max_h = 5.2 * inch
            if h > max_h:
                h = max_h
                w = h * iw / ih
            story.append(Image(os.path.join(HERE, path), width=w, height=h))
        elif kind == "tbl":
            caption, rows = payload
            ncols = max(len(r) for r in rows)
            data = [[mk(caption, capcell)] + [""] * (ncols - 1)]
            for ri, row in enumerate(rows):
                out = []
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ""
                    st = cellh if ri == 0 else (cell if ci == 0 else cellc)
                    out.append(mk(txt, st))
                data.append(out)
            weights = []
            for ci in range(ncols):
                mx = max(len(strip_markup(r[ci])) if ci < len(r) else 1 for r in rows)
                weights.append(max(6.0, min(mx, 46)) ** 0.85)
            tot = sum(weights)
            widths = [FRAME_W * w / tot for w in weights]
            t = Table(data, colWidths=widths, repeatRows=2, hAlign="CENTER")
            t.setStyle(TableStyle([
                ("SPAN", (0, 0), (-1, 0)),
                ("GRID", (0, 1), (-1, -1), 0.4, colors.HexColor("#9a9a9a")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#ececec")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (0, 0), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 2.4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
            ]))
            story.append(KeepTogether([Spacer(1, 6), t, Spacer(1, 10)]))
        elif kind == "table":
            rows = payload
            ncols = max(len(r) for r in rows)
            data = []
            for ri, row in enumerate(rows):
                out = []
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ""
                    st = cellh if ri == 0 else (cell if ci == 0 else cellc)
                    out.append(mk(txt, st))
                data.append(out)
            # column widths from content length
            weights = []
            for ci in range(ncols):
                mx = max(len(strip_markup(r[ci])) if ci < len(r) else 1 for r in rows)
                weights.append(max(6.0, min(mx, 46)) ** 0.85)
            tot = sum(weights)
            widths = [FRAME_W * w / tot for w in weights]
            t = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9a9a9a")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ececec")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 2.4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
            ]))
            if len(rows) <= 15:
                story.append(KeepTogether([Spacer(1, 3), t]))
            else:
                story.append(Spacer(1, 3))
                story.append(t)
            story.append(Spacer(1, 9))

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("DJSerif", 8.4)
        canvas.drawCentredString(letter[0] / 2, 0.62 * inch, str(doc.page))
        canvas.restoreState()

    doc = BaseDocTemplate(
        PDF_OUT, pagesize=letter,
        leftMargin=inch, rightMargin=inch, topMargin=0.95 * inch, bottomMargin=0.95 * inch,
        title="XAI Driven Modality Bias Detection and Mitigation in Multimodal Deep Learning Models",
        author="Nishant, I24AI028",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=on_page)])
    doc.build(story)
    print("wrote", PDF_OUT)


if __name__ == "__main__":
    text = open(SRC, encoding="utf-8").read()
    blocks = parse_blocks(text)
    build_docx(blocks)
    build_pdf(blocks)
